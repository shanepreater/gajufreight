"""Build the GajuFreight pitch deck for shippers and forwarders.

Run ``uv run python docs/pitch/build_deck.py`` after editing the slides below;
it rewrites ``gajufreight-pitch-deck.pptx`` beside this file. Colours come from
the brand guide (docs/brand/brand-guide.md) and type is Montserrat, so install
it to present (free from Google Fonts; otherwise PowerPoint substitutes a
system font). ``logo-horizontal.png`` is ``docs/brand/logo/logo-horizontal.svg``
rendered at 1600 px and cropped, because PowerPoint can't place SVG reliably.

Every claim here must match the design docs: fees (ADR 0010), what's public
on-chain (HLD §6.8) and the rollout (deployment plan §5).
"""

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.presentation import Presentation as Deck
from pptx.slide import Slide
from pptx.util import Emu, Inches, Length, Pt

HERE = Path(__file__).parent
OUTPUT = HERE / "gajufreight-pitch-deck.pptx"
LOGO = HERE / "logo-horizontal.png"

# Brand tokens (light theme). Signal Orange is never text on white: use ACCENT_TEXT.
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x4B, 0x55, 0x63)
LINE = RGBColor(0xD1, 0xD5, 0xDB)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE_ALT = RGBColor(0xF3, 0xF4, 0xF6)
ACCENT = RGBColor(0xFF, 0x57, 0x22)
ACCENT_TEXT = RGBColor(0xC2, 0x41, 0x0C)
ACCENT_SURFACE = RGBColor(0xFF, 0xF1, 0xEB)
FONT = "Montserrat"

WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_WIDTH = Emu(WIDTH - 2 * MARGIN)

# Fill in before sending the deck to a customer.
CONTACT = "Contact details to add"


@dataclass(frozen=True)
class Point:
    """A heading with a sentence of explanation, as used on card and list slides."""

    title: str
    body: str


def add_text(
    slide: Slide,
    left: Length,
    top: Length,
    width: Length,
    height: Length,
    text: str,
    size: int,
    *,
    bold: bool = False,
    color: RGBColor = INK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    """Place one text box; newlines in ``text`` become paragraphs."""
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(size * 0.4)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_box(
    slide: Slide,
    left: Length,
    top: Length,
    width: Length,
    height: Length,
    fill: RGBColor,
) -> None:
    """Draw a borderless filled rectangle (square corners, as the brand asks)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False


def new_slide(deck: Deck, eyebrow: str, title: str, notes: str) -> Slide:
    """Start a content slide: orange rule, eyebrow, headline, footer logo, notes."""
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_box(slide, MARGIN, Inches(0.55), Inches(0.9), Inches(0.08), ACCENT)
    add_text(
        slide,
        MARGIN,
        Inches(0.75),
        Inches(11),
        Inches(0.4),
        eyebrow.upper(),
        12,
        bold=True,
        color=ACCENT_TEXT,
    )
    add_text(
        slide, MARGIN, Inches(1.1), Inches(11.9), Inches(1.2), title, 30, bold=True
    )
    slide.shapes.add_picture(str(LOGO), MARGIN, Inches(6.85), height=Inches(0.3))
    number = len(deck.slides)
    add_text(
        slide,
        Emu(WIDTH - MARGIN - Inches(1)),
        Inches(6.82),
        Inches(1),
        Inches(0.4),
        str(number),
        11,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_cards(slide: Slide, points: list[Point], top: Length, height: Length) -> None:
    """Lay out points as equal cards in one row, each with a dark top rule."""
    gap = Inches(0.3)
    count = len(points)
    width = Emu(int((CONTENT_WIDTH - gap * (count - 1)) / count))
    for index, point in enumerate(points):
        left = Emu(MARGIN + index * (width + gap))
        add_box(slide, left, top, width, Inches(0.06), INK)
        add_text(
            slide,
            left,
            Emu(top + Inches(0.2)),
            width,
            Inches(0.9),
            point.title,
            16,
            bold=True,
        )
        add_text(
            slide,
            left,
            Emu(top + Inches(1.15)),
            width,
            Emu(height - Inches(1.15)),
            point.body,
            13,
            color=MUTED,
        )


def add_list(
    slide: Slide, points: list[Point], left: Length, top: Length, width: Length
) -> None:
    """Stack points down a column: bold heading, muted sentence beneath."""
    for index, point in enumerate(points):
        y = Emu(top + index * Inches(0.95))
        add_box(slide, left, Emu(y + Inches(0.12)), Inches(0.12), Inches(0.12), ACCENT)
        add_text(
            slide,
            Emu(left + Inches(0.3)),
            y,
            Emu(width - Inches(0.3)),
            Inches(0.4),
            point.title,
            16,
            bold=True,
        )
        add_text(
            slide,
            Emu(left + Inches(0.3)),
            Emu(y + Inches(0.38)),
            Emu(width - Inches(0.3)),
            Inches(0.55),
            point.body,
            12,
            color=MUTED,
        )


def title_slide(deck: Deck) -> None:
    """Slide 1: the promise in one line."""
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_box(slide, Emu(0), Emu(0), Inches(0.35), HEIGHT, ACCENT)
    slide.shapes.add_picture(str(LOGO), Inches(1.2), Inches(1.4), height=Inches(0.8))
    add_text(
        slide,
        Inches(1.2),
        Inches(2.9),
        Inches(11.5),
        Inches(1.4),
        "Freight paid as it moves",
        44,
        bold=True,
    )
    add_text(
        slide,
        Inches(1.2),
        Inches(4.2),
        Inches(10),
        Inches(1),
        "Escrow-backed shipping payments for shippers, forwarders and carriers",
        22,
        color=MUTED,
    )
    add_text(
        slide,
        Inches(1.2),
        Inches(6.3),
        Inches(10),
        Inches(0.5),
        "Pilot programme · 2026",
        14,
        bold=True,
        color=ACCENT_TEXT,
    )
    slide.notes_slide.notes_text_frame.text = (
        "GajuFreight agrees the price, locks the money and pays it out as the "
        "goods move, with proof at every step. Nobody has to trust anybody's "
        "spreadsheet, including ours."
    )


def problem_slide(deck: Deck) -> None:
    """Slide 2: what's wrong with freight payment today."""
    slide = new_slide(
        deck,
        "The problem",
        "Freight payment runs on trust that nobody can check",
        "Ask the room which of these they've lived through this quarter. "
        "The common thread: money and proof live in different places, so "
        "someone always carries the risk.",
    )
    add_cards(
        slide,
        [
            Point(
                "Shippers pay blind",
                "Paying up front, or on invoice, with little proof of where "
                "the goods are or whether milestones were met.",
            ),
            Point(
                "Forwarders wait",
                "Forwarders and carriers often fund the move themselves and "
                "wait weeks after delivery to be paid.",
            ),
            Point(
                "Status is scattered",
                "Updates arrive by email, portal and phone call, and each "
                "party holds a different version of the truth.",
            ),
            Point(
                "Disputes drag on",
                "With no shared record, a missing pallet becomes weeks of "
                "emails, withheld payments and sometimes lawyers.",
            ),
        ],
        Inches(2.6),
        Inches(3.6),
    )


def solution_slide(deck: Deck) -> None:
    """Slide 3: how GajuFreight works, in four steps."""
    slide = new_slide(
        deck,
        "How it works",
        "The deal, the money and the proof, in one shared record",
        "Every step is signed by the party taking it, in their own wallet. "
        "GajuFreight builds the transaction; it never holds keys or money.",
    )
    steps = [
        Point(
            "1  Agree",
            "The shipper invites forwarders to quote. Up to three counter "
            "offers, then the shipper accepts. The terms are fixed.",
        ),
        Point(
            "2  Fund",
            "One signature locks the agreed price in an escrow contract, on "
            "exactly the agreed terms.",
        ),
        Point(
            "3  Move",
            "Named port agents and carriers scan custody at each milestone. "
            "Each scan pays that milestone.",
        ),
        Point(
            "4  Settle",
            "Proven delivery releases the rest. Missed deadline: the shipper "
            "is refunded. Dispute: a panel decides.",
        ),
    ]
    add_cards(slide, steps, Inches(2.6), Inches(3.6))


def comparison_slide(deck: Deck) -> None:
    """Slide 4: side-by-side with the traditional approach."""
    slide = new_slide(
        deck,
        "Why switch",
        "Traditional freight payment and GajuFreight, side by side",
        "This is the slide to leave up during questions. Each GajuFreight cell "
        "is enforced by the contract, not by a promise in a terms page.",
    )
    rows = [
        ("", "Traditional", "GajuFreight"),
        (
            "Agreeing a price",
            "Emails, PDFs and phone calls",
            "Quotes and counters recorded; accepted terms can't change",
        ),
        (
            "Payment security",
            "Prepay and hope, or extend credit and chase",
            "Funds locked in escrow before anything moves",
        ),
        (
            "Getting paid",
            "On invoice terms, weeks after delivery",
            "At each confirmed milestone, within minutes",
        ),
        (
            "Proof of delivery",
            "Paper signatures and emailed photos",
            "Scan, photos and a delivery code, sealed with a fingerprint",
        ),
        (
            "Disputes",
            "Withheld payments, email chains, lawyers",
            "Panel of arbiters, with a fallback split agreed up front",
        ),
        (
            "Visibility",
            "Ask the forwarder",
            "Every party works from the same live record",
        ),
    ]
    shape = slide.shapes.add_table(
        len(rows), 3, MARGIN, Inches(2.3), CONTENT_WIDTH, Inches(4.2)
    )
    table = shape.table
    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(4.3)
    table.columns[2].width = Emu(int(CONTENT_WIDTH - Inches(6.9)))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            is_header = r == 0
            cell.fill.fore_color.rgb = (
                INK if is_header else ACCENT_SURFACE if c == 2 else SURFACE
            )
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(13 if is_header else 12)
            run.font.bold = is_header or c == 0
            run.font.color.rgb = SURFACE if is_header else INK if c != 1 else MUTED


def shippers_slide(deck: Deck) -> None:
    """Slide 5: what shippers gain."""
    slide = new_slide(
        deck,
        "For shippers",
        "Pay for proven progress, never for promises",
        "The shipper pays exactly the agreed price: the platform fee comes "
        "out of the forwarder's payouts, and refunds carry no fee.",
    )
    add_list(
        slide,
        [
            Point(
                "Your money is never in our hands",
                "Funds sit in a contract only the agreed rules can move. "
                "GajuFreight can't touch them.",
            ),
            Point(
                "Pay only as milestones are proven",
                "Each payment follows a custody scan by an attestor you agreed "
                "to, with evidence kept on file.",
            ),
            Point(
                "Refunds the contract guarantees",
                "If the deadline passes without delivery, the unpaid balance "
                "comes back to you. No chasing.",
            ),
            Point(
                "Comparable quotes, exact price",
                "Invited forwarders quote on the same consignment; you pay "
                "the price you accepted and no platform fee.",
            ),
        ],
        MARGIN,
        Inches(2.4),
        Inches(11.9),
    )


def forwarders_slide(deck: Deck) -> None:
    """Slide 6: what forwarders and carriers gain."""
    slide = new_slide(
        deck,
        "For forwarders and carriers",
        "Get paid as you deliver, not weeks later",
        "Fee: 1% of the shipping cost, taken from payouts, with a minimum of "
        "1 Gaju. The contract never takes more than 10% of the shipping cost, "
        "even for a tiny shipment. The rate is fixed when the quote is "
        "requested, so it can't change mid-shipment.",
    )
    add_list(
        slide,
        [
            Point(
                "Know the money is there before you move",
                "The shipper funds the escrow when they book, so you never "
                "carry the credit risk.",
            ),
            Point(
                "Paid at every milestone",
                "Each confirmed scan-in pays its milestone straight to you, "
                "improving cash flow on long moves.",
            ),
            Point(
                "Subcontract legs on the same terms",
                "Book each carrier through GajuFreight too: they're paid at "
                "handover, from an escrow you fund.",
            ),
            Point(
                "A dispute can't freeze everything",
                "Milestones already paid stay paid. Only the unpaid balance "
                "goes to the panel.",
            ),
        ],
        MARGIN,
        Inches(2.4),
        Inches(11.9),
    )


def trust_slide(deck: Deck) -> None:
    """Slide 7: why the system can be trusted."""
    slide = new_slide(
        deck,
        "Built-in trust",
        "Nobody has to trust us, or each other",
        "Gajumaru is a public blockchain. The contracts are small, open "
        "source and will be externally audited before any real money flows.",
    )
    add_cards(
        slide,
        [
            Point(
                "We never hold keys or funds",
                "Every action is signed in the party's own wallet. Our "
                "servers build transactions; they can't sign them.",
            ),
            Point(
                "Named witnesses per shipment",
                "Only the attestors agreed for this shipment can record "
                "custody, and no one attests their own payment.",
            ),
            Point(
                "Tamper-evident evidence",
                "Scans, photos and documents are stored encrypted; their "
                "fingerprint on the ledger proves they haven't changed.",
            ),
            Point(
                "No single party can block payment",
                "Delivery codes, challenge windows, deadlines and a panel "
                "fallback mean money always has a way out.",
            ),
        ],
        Inches(2.6),
        Inches(3.6),
    )


def privacy_slide(deck: Deck) -> None:
    """Slide 8: what's public and what stays private, stated plainly."""
    slide = new_slide(
        deck,
        "Privacy",
        "What's on the public ledger, and what stays private",
        "Say this upfront: agreed terms are public so both sides can rely on "
        "them later. A competitor reading the ledger could see lanes, volumes "
        "and prices, but never who the parties are or what the documents say.",
    )
    half = Emu(int((CONTENT_WIDTH - Inches(0.4)) / 2))
    columns = [
        (
            MARGIN,
            SURFACE_ALT,
            "Public on the ledger",
            "Agreed price and milestone schedule\n"
            "Route, as port codes (for example CNSHA → NLRTM)\n"
            "Cargo sizes, weights and deliver-by date\n"
            "Platform fees paid",
        ),
        (
            Emu(MARGIN + half + Inches(0.4)),
            ACCENT_SURFACE,
            "Private to those involved",
            "Company names, contacts and addresses\n"
            "The consignee\n"
            "Documents, photos and delivery evidence\n"
            "Leg prices, shown only to the forwarder and that carrier",
        ),
    ]
    for left, fill, heading, body in columns:
        add_box(slide, left, Inches(2.4), half, Inches(3.9), fill)
        add_text(
            slide,
            Emu(left + Inches(0.35)),
            Inches(2.65),
            Emu(half - Inches(0.7)),
            Inches(0.5),
            heading,
            18,
            bold=True,
        )
        add_text(
            slide,
            Emu(left + Inches(0.35)),
            Inches(3.35),
            Emu(half - Inches(0.7)),
            Inches(2.8),
            body,
            14,
        )


def cost_slide(deck: Deck) -> None:
    """Slide 9: what it costs each party."""
    slide = new_slide(
        deck,
        "Pricing",
        "Simple fees, fixed before you agree",
        "Network fees are about 0.0002 Gaju per action, so a whole shipment "
        "with two legs costs about 0.0044 Gaju in network fees, shared "
        "between the parties. Prices and payments are in Gaju (木) for now; "
        "other currencies are out of scope for the pilot. The fee starts at "
        "1% with a 1 Gaju minimum; both are set by GajuFreight's admin vote "
        "and can change, but a shipment keeps the settings in force when its "
        "quotes were requested. The 10% limit is fixed in the contract: no "
        "vote can raise the rate above it, and it overrides the minimum.",
    )
    add_cards(
        slide,
        [
            Point(
                "Shippers",
                "Pay the agreed price, and nothing more. Refunds and your share "
                "of a dispute split carry no fee.",
            ),
            Point(
                "Forwarders",
                "1% of the shipping cost (minimum 1 Gaju), fixed when the "
                "quote is requested. Never more than 10%: a limit built into "
                "the contract.",
            ),
            Point(
                "Network fees",
                "A tiny charge per signed action, paid by whoever signs it: a "
                "fraction of a Gaju for a whole shipment.",
            ),
        ],
        Inches(2.6),
        Inches(3.4),
    )


def pilot_slide(deck: Deck) -> None:
    """Slide 10: the ask."""
    slide = new_slide(
        deck,
        "Next steps",
        "Join the pilot",
        "We're looking for a handful of shippers and forwarders with real "
        "lanes to shape the product. The testnet stage uses no real money.",
    )
    add_cards(
        slide,
        [
            Point(
                "Testnet alpha",
                "Run real lanes end to end with test funds. No money at risk; "
                "your feedback shapes the screens.",
            ),
            Point(
                "Mainnet pilot",
                "Real payments with a cap on shipment value, after an external "
                "audit of the contracts.",
            ),
            Point(
                "What we need from you",
                "A lane or two, a GajuDesk or GajuMobile wallet, and regular, "
                "honest feedback.",
            ),
        ],
        Inches(2.4),
        Inches(2.9),
    )
    add_box(slide, MARGIN, Inches(5.6), CONTENT_WIDTH, Inches(0.9), INK)
    add_text(
        slide,
        Emu(MARGIN + Inches(0.35)),
        Inches(5.8),
        Emu(CONTENT_WIDTH - Inches(0.7)),
        Inches(0.6),
        CONTACT,
        18,
        bold=True,
        color=SURFACE,
    )


def build() -> None:
    """Write the deck, one function per slide, in presenting order."""
    deck = Presentation()
    deck.slide_width = WIDTH
    deck.slide_height = HEIGHT
    for slide in (
        title_slide,
        problem_slide,
        solution_slide,
        comparison_slide,
        shippers_slide,
        forwarders_slide,
        trust_slide,
        privacy_slide,
        cost_slide,
        pilot_slide,
    ):
        slide(deck)
    deck.save(str(OUTPUT))


if __name__ == "__main__":
    build()
